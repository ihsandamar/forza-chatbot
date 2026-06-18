import logging
from io import BytesIO
from typing import Any, Dict, List

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor

from src.graphs.report_agent.subagents.export_v1.exporters.base import BaseExporter
from src.graphs.report_agent.subagents.export_v1.exporters.registry import ExporterRegistry
from src.graphs.report_agent.subagents.export_v1.models import ExportSpec

logger = logging.getLogger(__name__)

_HEADER_RGB = RGBColor(0x1F, 0x4E, 0x79)
_HEADER_HEX = "1F4E79"


@ExporterRegistry.register
class DocxExporter(BaseExporter):
    FORMAT = "docx"

    @property
    def content_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def export(
        self,
        results: List[Dict[str, Any]],
        spec: ExportSpec,
        chart_figure=None,
        insight_text: str = "",
    ) -> BytesIO:
        doc = Document()

        # Heading 1 with navy colour
        heading = doc.add_heading(spec.title, level=1)
        if heading.runs:
            heading.runs[0].font.color.rgb = _HEADER_RGB

        # Optional insight paragraph
        if spec.include_insight and insight_text:
            doc.add_paragraph(insight_text)

        # Optional chart image — python-docx accepts a BytesIO directly
        if chart_figure is not None:
            try:
                img_buf = BytesIO()
                chart_figure.savefig(img_buf, format="png", dpi=150, bbox_inches="tight")
                img_buf.seek(0)
                doc.add_picture(img_buf, width=Inches(5.5))
            except Exception as exc:
                logger.warning(f"[DOCX] Chart embedding failed: {exc}")

        # Data table or empty notice
        if results:
            headers = list(results[0].keys())
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"

            hdr_cells = table.rows[0].cells
            for i, h in enumerate(headers):
                cell = hdr_cells[i]
                cell.text = str(h)
                run = cell.paragraphs[0].runs[0]
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                _set_cell_bg(cell, _HEADER_HEX)

            for row in results:
                cells = table.add_row().cells
                for i, h in enumerate(headers):
                    cells[i].text = str(row.get(h, ""))
        else:
            doc.add_paragraph("Veri bulunamadı.")

        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf


def _set_cell_bg(cell, hex_color: str) -> None:
    """Apply a solid background fill to a table cell via OOXML."""
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)
