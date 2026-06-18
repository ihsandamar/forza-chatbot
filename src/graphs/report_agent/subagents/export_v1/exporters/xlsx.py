import logging
import os
import tempfile
from io import BytesIO
from typing import Any, Dict, List

import openpyxl
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill

from src.graphs.report_agent.subagents.export_v1.exporters.base import BaseExporter
from src.graphs.report_agent.subagents.export_v1.exporters.registry import ExporterRegistry
from src.graphs.report_agent.subagents.export_v1.models import ExportSpec

logger = logging.getLogger(__name__)

_HEADER_COLOR = "1F4E79"
_ALT_ROW_COLOR = "EBF5FB"


@ExporterRegistry.register
class XlsxExporter(BaseExporter):
    FORMAT = "xlsx"

    @property
    def content_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def export(
        self,
        results: List[Dict[str, Any]],
        spec: ExportSpec,
        chart_figure=None,
        insight_text: str = "",
    ) -> BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = spec.sheet_name

        ws.cell(row=1, column=1, value=spec.title).font = Font(bold=True, size=14)
        ws.row_dimensions[1].height = 24

        if not results:
            ws.cell(row=3, column=1, value="Veri bulunamadı.")
            buf = BytesIO()
            wb.save(buf)
            return buf

        headers = list(results[0].keys())
        header_fill = PatternFill("solid", fgColor=_HEADER_COLOR)
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        alt_fill = PatternFill("solid", fgColor=_ALT_ROW_COLOR)
        for row_idx, row in enumerate(results, 4):
            for col_idx, key in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=row.get(key))
                if (row_idx - 4) % 2 == 1:
                    cell.fill = alt_fill

        for col_idx, header in enumerate(headers, 1):
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            max_len = max(
                len(str(header)),
                *[len(str(row.get(header, ""))) for row in results],
            )
            ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

        _tmp_chart_path = None

        if chart_figure is not None:
            ws_chart = wb.create_sheet("Grafik")
            ws_chart.cell(row=1, column=1, value=spec.title).font = Font(bold=True, size=14)
            try:
                with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                    _tmp_chart_path = tmp.name
                chart_figure.savefig(_tmp_chart_path, dpi=150, bbox_inches="tight")
                img = XlImage(_tmp_chart_path)
                img.anchor = "B3"
                ws_chart.add_image(img)
            except Exception as exc:
                logger.warning(f"[XLSX] Chart embedding failed: {exc}")
                if _tmp_chart_path and os.path.exists(_tmp_chart_path):
                    os.unlink(_tmp_chart_path)
                _tmp_chart_path = None

        buf = BytesIO()
        try:
            wb.save(buf)
        finally:
            if _tmp_chart_path and os.path.exists(_tmp_chart_path):
                os.unlink(_tmp_chart_path)
        return buf
