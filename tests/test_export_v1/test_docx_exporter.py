import pytest
from io import BytesIO

import src.graphs.report_agent.subagents.export_v1.exporters.docx  # noqa: F401
from src.graphs.report_agent.subagents.export_v1.exporters.registry import ExporterRegistry
from src.graphs.report_agent.subagents.export_v1.models import ExportSpec
from docx import Document

SAMPLE_RESULTS = [
    {"BranchName": "Istanbul", "RowAmount": 50000},
    {"BranchName": "Ankara", "RowAmount": 35000},
]


def test_docx_produces_valid_document():
    exporter = ExporterRegistry.get("docx")
    spec = ExportSpec(format="docx", title="Rapor — 2026-06-18")
    buf = exporter.export(SAMPLE_RESULTS, spec)
    buf.seek(0)
    doc = Document(buf)
    assert len(doc.paragraphs) > 0


def test_docx_title_in_document():
    exporter = ExporterRegistry.get("docx")
    spec = ExportSpec(format="docx", title="Rapor — 2026-06-18")
    buf = exporter.export(SAMPLE_RESULTS, spec)
    buf.seek(0)
    doc = Document(buf)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any("Rapor" in h for h in headings)


def test_docx_table_has_correct_data():
    exporter = ExporterRegistry.get("docx")
    spec = ExportSpec(format="docx", title="Rapor")
    buf = exporter.export(SAMPLE_RESULTS, spec)
    buf.seek(0)
    doc = Document(buf)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    # header row
    assert table.rows[0].cells[0].text == "BranchName"
    assert table.rows[1].cells[0].text == "Istanbul"


def test_docx_empty_results():
    exporter = ExporterRegistry.get("docx")
    spec = ExportSpec(format="docx", title="Rapor")
    buf = exporter.export([], spec)
    buf.seek(0)
    doc = Document(buf)
    texts = [p.text for p in doc.paragraphs]
    assert any("bulunamadı" in t for t in texts)


def test_docx_insight_added_when_flag_set():
    exporter = ExporterRegistry.get("docx")
    spec = ExportSpec(format="docx", title="Rapor", include_insight=True)
    buf = exporter.export(SAMPLE_RESULTS, spec, insight_text="Toplam satış arttı.")
    buf.seek(0)
    doc = Document(buf)
    texts = [p.text for p in doc.paragraphs]
    assert any("Toplam satış arttı." in t for t in texts)


def test_docx_insight_not_added_when_flag_false():
    exporter = ExporterRegistry.get("docx")
    spec = ExportSpec(format="docx", title="Rapor", include_insight=False)
    buf = exporter.export(SAMPLE_RESULTS, spec, insight_text="Gizli bilgi.")
    buf.seek(0)
    doc = Document(buf)
    texts = [p.text for p in doc.paragraphs]
    assert not any("Gizli bilgi." in t for t in texts)


def test_docx_chart_embedded_when_figure_given():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots()
    ax.bar(["A", "B"], [1, 2])
    exporter = ExporterRegistry.get("docx")
    spec = ExportSpec(format="docx", title="Rapor", include_chart=True)
    buf = exporter.export(SAMPLE_RESULTS, spec, chart_figure=fig)
    buf.seek(0)
    doc = Document(buf)
    # When a chart is embedded, there should be inline shapes (images) in the doc
    assert doc.inline_shapes is not None
    assert len(doc.inline_shapes) > 0
    plt.close("all")


def test_docx_content_type():
    exporter = ExporterRegistry.get("docx")
    assert exporter.content_type == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
