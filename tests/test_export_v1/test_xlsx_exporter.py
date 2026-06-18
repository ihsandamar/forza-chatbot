import openpyxl
import pytest
from io import BytesIO

import src.graphs.report_agent.subagents.export_v1.exporters.xlsx  # noqa: F401
from src.graphs.report_agent.subagents.export_v1.exporters.registry import ExporterRegistry
from src.graphs.report_agent.subagents.export_v1.models import ExportSpec

SAMPLE_RESULTS = [
    {"BranchName": "Istanbul", "RowAmount": 50000},
    {"BranchName": "Ankara", "RowAmount": 35000},
]


def test_xlsx_produces_valid_workbook():
    exporter = ExporterRegistry.get("xlsx")
    spec = ExportSpec(format="xlsx", title="Rapor — 2026-06-18")
    buf = exporter.export(SAMPLE_RESULTS, spec)
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)
    assert "Veri" in wb.sheetnames


def test_xlsx_data_rows_correct():
    exporter = ExporterRegistry.get("xlsx")
    spec = ExportSpec(format="xlsx", title="Rapor — 2026-06-18")
    buf = exporter.export(SAMPLE_RESULTS, spec)
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)
    ws = wb["Veri"]
    # row 3 = header, row 4 = first data row
    assert ws.cell(row=3, column=1).value == "BranchName"
    assert ws.cell(row=4, column=1).value == "Istanbul"
    assert ws.cell(row=5, column=1).value == "Ankara"


def test_xlsx_chart_sheet_created_when_figure_given():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots()
    ax.bar(["A", "B"], [1, 2])
    exporter = ExporterRegistry.get("xlsx")
    spec = ExportSpec(format="xlsx", title="Rapor", include_chart=True)
    buf = exporter.export(SAMPLE_RESULTS, spec, chart_figure=fig)
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)
    assert "Grafik" in wb.sheetnames
    plt.close("all")


def test_xlsx_no_chart_sheet_when_no_figure():
    exporter = ExporterRegistry.get("xlsx")
    spec = ExportSpec(format="xlsx", title="Rapor")
    buf = exporter.export(SAMPLE_RESULTS, spec, chart_figure=None)
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)
    assert "Grafik" not in wb.sheetnames
